from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Proyectos\Restaurante\output\Guia_de_uso_RestoManager.docx"
NAVY, RED, PALE, GRID = "17324D", "E63946", "EEF3F7", "D9E0E7"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.67), Inches(.65)
section.left_margin = section.right_margin = Inches(.72)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.2)
styles["Normal"].font.color.rgb = RGBColor.from_string("243244")
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 30), ("Heading 1", 19), ("Heading 2", 13)):
    s = styles[name]
    s.font.name = "Aptos Display" if name != "Heading 2" else "Aptos"
    s.font.size, s.font.bold = Pt(size), True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    s.paragraph_format.space_after = Pt(5)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("RESTOMANAGER  /  GUÍA DE OPERACIÓN")
r.font.name, r.font.size, r.font.bold = "Aptos", Pt(8), True
r.font.color.rgb = RGBColor.from_string("65758A")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Manual para usuarios   ·   ")
r.font.name, r.font.size = "Aptos", Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def make_table(headers, rows, widths, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "5")
        e.set(qn("w:color"), GRID)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "95"), ("start", "115"), ("bottom", "95"), ("end", "115")):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), value)
        e.set(qn("w:type"), "dxa")
        margins.append(e)
    t._tbl.tblPr.append(margins)
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = text
        shade(c, NAVY)
        for run in c.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
    tr = OxmlElement("w:tblHeader")
    tr.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(tr)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            c = cells[i]
            c.text = str(text)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2:
                shade(c, PALE)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name, run.font.size = "Aptos", Pt(font_size)
        for i, width in enumerate(widths):
            cells[i].width = Inches(width)
    for i, width in enumerate(widths):
        t.rows[0].cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def heading(text, level=1):
    doc.add_heading(text, level=level)

def para(text):
    doc.add_paragraph(text)

def bullets(items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(text)

def steps(items):
    for text in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(text)

# Portada y recorrido general
p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(42), Pt(6)
r = p.add_run("GUÍA DE USO")
r.font.name, r.font.size, r.font.bold = "Aptos", Pt(11), True
r.font.color.rgb = RGBColor.from_string(RED)
p = doc.add_paragraph(style="Title")
p.add_run("Sistema de gestión\npara restaurantes")
p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(24)
r = p.add_run("Procedimientos para caja, salón, cocina, reparto y administración")
r.font.name, r.font.size = "Aptos", Pt(14)
r.font.color.rgb = RGBColor.from_string("536579")
para("Esta guía explica cómo iniciar sesión, registrar pedidos y completar el servicio hasta el pago y el cierre de caja. Cada persona debe usar su propia cuenta para que el sistema asigne correctamente las comandas, los cobros y las entregas.")
heading("El recorrido de un pedido", 2)
make_table(["Etapa", "Qué ocurre"], [
    ("1. Pedido", "Mesero o cajero arma una comanda en el POS; el cliente también puede pedir desde la carta QR."),
    ("2. Preparación", "Cocina inicia cada producto y lo marca listo."),
    ("3. Entrega", "El pedido listo se sirve en mesa o se asigna a un repartidor."),
    ("4. Pago", "Caja registra el pago, actualiza el turno y libera la mesa al completar el cobro."),
], [1.35, 5.75], 9.5)
para("Uso recomendado: conecte los equipos de operación a la misma red Wi-Fi del restaurante para que cocina y notificaciones se actualicen en tiempo real.")
para("El enlace temporal de Cloudflare sirve para acceder desde fuera de la red, pero no mantiene los eventos de cocina en tiempo real. Para una operación remota continua se requiere un túnel Cloudflare permanente.")

# Acceso y roles
doc.add_page_break()
heading("1 Acceso y responsabilidades")
heading("Cómo iniciar sesión", 2)
steps([
    "Abra el enlace que le entregó el administrador. En el restaurante puede usar la dirección local del equipo servidor; desde el mismo Wi-Fi puede usar su dirección IP local.",
    "Elija PIN Rápido e ingrese su PIN de cuatro dígitos, o seleccione Email / Clave y escriba sus credenciales.",
    "Si el acceso falla, confirme que la cuenta esté activa y que el PIN o la contraseña sean correctos. Después de varios intentos fallidos con PIN, espere antes de volver a intentar.",
    "Al cambiar de empleado, cierre la sesión o use Cambiar con PIN para atribuir las operaciones a la persona correcta.",
])
heading("Qué puede hacer cada rol", 2)
make_table(["Rol", "Responsabilidades principales", "Acceso habitual"], [
    ("Administrador", "Configura el negocio, usuarios, menú, mesas, caja, inventario, delivery, finanzas y reportes.", "Panel y módulos administrativos"),
    ("Cajero", "Abre su turno, cobra pedidos, registra movimientos y consulta reportes y clientes.", "POS, caja, delivery y finanzas permitidas"),
    ("Mesero", "Crea y actualiza comandas, asigna mesa y comensales, y envía pedidos a cocina.", "POS"),
    ("Cocinero", "Ve las comandas pendientes, inicia productos y marca productos o pedidos listos.", "Cocina KDS"),
    ("Repartidor", "Consulta sus entregas asignadas y confirma que un pedido enviado fue entregado.", "Delivery asignado"),
], [1.05, 3.6, 2.45], 8.5)
para("El sistema verifica el rol en cada operación. Si un módulo no aparece, pida al administrador que confirme su rol; no utilice la cuenta de otra persona.")

# POS y cocina
doc.add_page_break()
heading("2 Registrar un pedido")
heading("Pedido de mesa, mostrador o domicilio", 2)
steps([
    "Abra Terminal POS. Seleccione En Mesa, Mostrador o Delivery.",
    "Para una mesa, elija la sala y la mesa. Verifique que sea la correcta antes de agregar productos.",
    "Busque una categoría o escriba el nombre del plato. Toque el producto para agregarlo; elija modificadores cuando se soliciten y anote instrucciones importantes.",
    "En pedidos de mesa, asigne los productos a sus comensales cuando necesite separar cuentas. Revise cantidades, notas y total en la comanda.",
    "Pulse Enviar a cocina una sola vez y espere el mensaje de confirmación. Para añadir productos después, agregue únicamente los nuevos y envíelos cuando estén listos.",
])
heading("Revisar una comanda en cocina", 2)
steps([
    "Abra Cocina KDS y localice la comanda por su número, mesa o tipo de pedido.",
    "Pulse Iniciar en cada producto cuando comience su preparación. Las instrucciones y modificadores aparecen junto al producto.",
    "Pulse Listo al terminar cada producto. Cuando toda la comanda esté preparada, use Marcar todo listo si corresponde.",
    "Para una comanda de mesa, pulse Entregado cuando el plato ya fue servido. La cocina deja de mostrar los productos servidos.",
])
para("Marcar listo o servido no registra el pago. La comanda permanece pendiente de cobro hasta que caja complete la cuenta.")
heading("Si se envió algo por error", 2)
para("Avise de inmediato al responsable de caja o al administrador. Evite volver a enviar repetidamente la misma comanda; el responsable revisará el pedido antes de corregirlo o cancelarlo.")

# Caja y reparto
doc.add_page_break()
heading("3 Abrir caja cobrar y cerrar")
heading("Abrir un turno", 2)
steps([
    "Entre a Caja & Turnos con su usuario de cajero o administrador.",
    "Escriba el efectivo inicial que recibe físicamente y pulse Abrir caja. Cada cajero debe abrir su propio turno.",
])
heading("Cobrar una comanda", 2)
steps([
    "Busque el pedido en Pedidos pendientes y ábralo.",
    "Elija cobrar el total o la cuenta de un comensal, si la mesa separará pagos.",
    "Revise el saldo y la propina. Seleccione el medio de pago; si se divide entre medios, indique cada monto.",
    "Confirme que los pagos sumen exactamente el total a cobrar y registre el pago una sola vez.",
    "Cuando el saldo completo queda pagado, el pedido se cierra y su mesa vuelve a estar disponible.",
])
para("Los medios disponibles incluyen efectivo, tarjetas, transferencia y cuenta corriente cuando esté habilitada. El inventario de recetas se descuenta al completar la venta.")
heading("Registrar movimientos y cerrar el turno", 2)
steps([
    "Registre cada entrada o salida manual de efectivo con el monto y un concepto claro.",
    "Al terminar el turno, cuente el efectivo de caja. Abra Cerrar turno e ingrese el efectivo real contado.",
    "Revise la diferencia que muestra el sistema, agregue observaciones si hace falta y confirme el cierre.",
])
para("El arqueo compara el efectivo contado con el saldo inicial, los cobros en efectivo y los movimientos manuales. Los pagos por tarjeta o transferencia no forman parte del efectivo físico.")
heading("4 Preparar y completar una entrega", 1)
steps([
    "El administrador o cajero abre Delivery, revisa la dirección y asigna un repartidor activo.",
    "El responsable de despacho avanza el pedido por su preparación y lo marca Enviado cuando sale del restaurante.",
    "El repartidor inicia sesión, abre Delivery y consulta los pedidos asignados a su cuenta.",
    "Después de entregar al cliente, el repartidor pulsa Confirmar entrega. No marque entregado antes de completar la entrega.",
])
para("La entrega y el pago son controles separados. Confirmar que llegó al cliente no cierra la cuenta ni reemplaza el cobro en caja.")

# QR y tareas administrativas
doc.add_page_break()
heading("5 Carta QR y pedidos del cliente")
steps([
    "El administrador abre Carta QR y pulsa Imprimir para generar un código por mesa.",
    "Coloque cada código en la mesa que indica su rótulo. Si cambió la dirección pública del sistema, vuelva a imprimir los códigos.",
    "El cliente escanea el código, selecciona platos y cantidades y abre su pedido.",
    "El cliente elige Mesa, Retiro o Domicilio; escribe nombre y teléfono, e indica la mesa o dirección que corresponda.",
    "Al confirmar, guarda el enlace de seguimiento para consultar el estado del pedido.",
])
para("Un QR de mesa debe apuntar al enlace público correcto y a la mesa correspondiente. Para probar la carta, confirme el pedido con el responsable antes de hacerlo, porque un pedido confirmado ingresa a la operación.")
heading("6 Tareas administrativas", 1)
make_table(["Módulo", "Uso"], [
    ("Menú y catálogo", "Crear categorías y platos, precios, disponibilidad, modificadores y favoritos."),
    ("Mesas y salas", "Organizar áreas y mesas, revisar estados, trasladar mesas y administrar reservaciones."),
    ("Inventario", "Crear ingredientes, registrar compras, mermas y conteos, y asociar recetas a platos."),
    ("Finanzas", "Mantener clientes y proveedores, cuentas corrientes, gastos y categorías de gastos."),
    ("Empleados", "Crear cuentas, asignar roles, configurar PIN y activar o desactivar usuarios."),
    ("Reportes", "Consultar ventas, productos, mesas, existencias críticas y gastos; exportar a Excel."),
    ("Configuración", "Editar datos del restaurante, marca lateral, moneda, impuestos, propina y alertas de cocina."),
], [1.55, 5.55], 9)
heading("Personalizar la marca lateral", 2)
para("En Configuración, dentro de Datos del Establecimiento, complete Nombre del encabezado y Texto secundario en la sección Marca del sistema. Pulse Guardar Configuración. Los cambios aparecen en el encabezado del menú lateral; el nombre comercial del restaurante es un campo diferente.")

# Lista diaria y ayuda
doc.add_page_break()
heading("7 Lista de revisión diaria")
heading("Antes de abrir", 2)
bullets([
    "Verifique que cada empleado tenga cuenta activa y el rol adecuado.",
    "Abra el turno de caja y registre el efectivo inicial real.",
    "Confirme que los platos disponibles, precios, modificadores y niveles de inventario estén correctos.",
    "Encienda la pantalla de cocina y compruebe que haya conexión con el sistema.",
])
heading("Al terminar", 2)
bullets([
    "Confirme que los pedidos servidos o entregados estén cobrados o documentados para seguimiento.",
    "Revise ingresos y egresos manuales de caja; cuente el efectivo.",
    "Cierre el turno con el monto contado y deje anotada cualquier diferencia.",
    "Registre mermas o conteos de inventario que deban actualizarse.",
])
heading("Solución rápida de problemas", 2)
make_table(["Situación", "Qué revisar"], [
    ("No puedo ingresar", "Use el método correcto, confirme PIN de cuatro dígitos o correo y contraseña, y pida validar que la cuenta siga activa."),
    ("No veo un módulo", "El acceso depende del rol. Pida al administrador revisar los permisos de su cuenta."),
    ("Un pedido no aparece en cocina", "Confirme el mensaje de envío del POS y busque por número de comanda. No pulse enviar repetidamente."),
    ("Un pago no se refleja", "Revise el pedido y el turno antes de intentar cobrar otra vez; confirme el medio y monto registrados."),
    ("El QR no abre", "Compruebe la conexión y la dirección del QR. Los enlaces temporales de Cloudflare cambian al reiniciar el túnel."),
    ("Cocina no se actualiza", "Use la red Wi-Fi local del restaurante. El túnel rápido de Cloudflare no admite eventos en tiempo real."),
], [1.45, 5.65], 8.6)
heading("Buenas prácticas", 2)
bullets([
    "No comparta PIN ni contraseña; cada operación debe quedar a nombre de quien la realizó.",
    "Espere la confirmación antes de repetir un envío o un cobro.",
    "Ante un pedido, pago o inventario dudoso, consulte primero el registro actual y avise al administrador.",
])
para("El acceso público depende de que el computador del restaurante, la base de datos y el túnel estén funcionando. Para uso permanente desde Internet, el administrador debe configurar un túnel con nombre en Cloudflare.")

doc.core_properties.title = "Guía de uso del sistema de gestión para restaurantes"
doc.core_properties.subject = "Procedimientos para usuarios de RestoManager"
doc.core_properties.author = "RestoManager"
doc.save(OUT)
print(OUT)
